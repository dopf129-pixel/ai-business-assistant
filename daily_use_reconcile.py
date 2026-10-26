"""Local-only CLI for comparing an Ozon accrual XLSX with a saved snapshot.

The command never calls Ozon or the bot. It reads one explicit report profile,
projects only safe totals, and writes JSON to stdout or a caller-selected path.
"""

import argparse
from datetime import date, datetime
import json
from pathlib import Path
import sys

sys.path.insert(0, str(Path(__file__).resolve().parent / "app"))

from services.daily_use_reconciliation_service import DailyUseReconciliationService


def _read_json(path):
    with Path(path).open("r", encoding="utf-8-sig") as stream:
        return json.load(stream)


def _write_json(value, output_path=None):
    payload = json.dumps(value, ensure_ascii=False, allow_nan=False, indent=2)
    if output_path:
        try:
            with Path(output_path).open("x", encoding="utf-8", newline="\n") as stream:
                stream.write(payload + "\n")
        except FileExistsError:
            print(json.dumps({"error": True, "code": "DAILY_USE_OUTPUT_ALREADY_EXISTS"}))
            return False
    else:
        print(payload)
    return True


def _read_accrual_workbook(path, service):
    try:
        from openpyxl import load_workbook
    except ImportError:
        return None, service._error("DAILY_USE_XLSX_READER_UNAVAILABLE")
    try:
        if Path(path).stat().st_size > 100 * 1024 * 1024:
            return None, service._error("DAILY_USE_XLSX_SIZE_LIMIT_EXCEEDED")
    except OSError:
        return None, service._error("DAILY_USE_XLSX_READ_FAILED")
    workbook = None
    try:
        workbook = load_workbook(path, read_only=True, data_only=True)
        if service.SHEET_NAME not in workbook.sheetnames:
            return None, service._error("DAILY_USE_REPORT_PROFILE_SHEET_MISMATCH")
        sheet = workbook[service.SHEET_NAME]
        sheet.reset_dimensions()
        header_values = next(
            sheet.iter_rows(
                min_row=service.HEADER_ROW, max_row=service.HEADER_ROW,
                values_only=True,
            ),
            None,
        )
        if header_values is None:
            return None, service._error("DAILY_USE_REPORT_PROFILE_COLUMNS_MISSING")
        headers = list(header_values)
        while headers and headers[-1] is None:
            headers.pop()
        if len(headers) > service.MAX_REPORT_COLUMNS:
            return None, service._error("DAILY_USE_REPORT_PROFILE_COLUMNS_INVALID")
        rows = []
        for values in sheet.iter_rows(
            min_row=service.HEADER_ROW + 1, values_only=True
        ):
            row = list(values[:len(headers)])
            if not any(value is not None for value in row):
                continue
            if len(rows) >= service.MAX_REPORT_ROWS:
                return None, service._error("DAILY_USE_REPORT_PROFILE_ROWS_INVALID")
            while len(row) < len(headers):
                row.append(None)
            rows.append(row)
        return {"sheet_name": sheet.title, "header_row": service.HEADER_ROW, "headers": headers, "rows": rows}, None
    except Exception:
        # Workbook exceptions may contain sensitive cell values. Return a stable
        # code rather than forwarding library messages or paths to the output.
        return None, service._error("DAILY_USE_XLSX_READ_FAILED")
    finally:
        if workbook is not None:
            workbook.close()


def _build_parser():
    parser = argparse.ArgumentParser(description="Offline report/snapshot comparison")
    commands = parser.add_subparsers(dest="command", required=True)

    reference = commands.add_parser("reference", help="Normalize one accrual XLSX")
    reference.add_argument("--profile", required=True)
    reference.add_argument("--report", required=True)
    reference.add_argument("--period-from", required=True)
    reference.add_argument("--period-to", required=True)
    reference.add_argument("--source-period-from", required=True)
    reference.add_argument("--source-period-to", required=True)
    reference.add_argument("--scope-alias", required=True)
    reference.add_argument("--confirm-full-period", action="store_true")
    reference.add_argument("--sku-filter")
    reference.add_argument("--sku-alias")
    reference.add_argument("--output")

    snapshot = commands.add_parser("snapshot", help="Project the saved Period Profit JSON export")
    snapshot.add_argument("--export", required=True)
    snapshot.add_argument("--scope-alias", required=True)
    snapshot.add_argument("--revision")
    snapshot.add_argument("--output")

    compare = commands.add_parser("compare", help="Compare normalized reference and snapshot")
    compare.add_argument("--reference", required=True)
    compare.add_argument("--snapshot", required=True)
    compare.add_argument("--output")
    return parser


def main(argv=None):
    args = _build_parser().parse_args(argv)
    service = DailyUseReconciliationService()
    if args.command == "reference":
        if args.profile != service.PROFILE_ID:
            result = service._error("DAILY_USE_REPORT_PROFILE_UNSUPPORTED")
        else:
            table, error = _read_accrual_workbook(args.report, service)
            result = error or service.normalize_accrual_table(
                **table,
                period_from=args.period_from,
                period_to=args.period_to,
                source_period_from=args.source_period_from,
                source_period_to=args.source_period_to,
                scope_alias=args.scope_alias,
                coverage_assertion=(
                    "USER_CONFIRMED_FULL_PERIOD" if args.confirm_full_period else None
                ),
                sku_filter_value=args.sku_filter,
                sku_dimension_alias=args.sku_alias,
            )
        written = _write_json(result, args.output)
        return 1 if result.get("error") or not written else 0

    if args.command == "snapshot":
        try:
            document = _read_json(args.export)
        except (OSError, ValueError, UnicodeError):
            result = service._error("DAILY_USE_CANONICAL_EXPORT_READ_FAILED")
        else:
            revision = args.revision
            if revision is None and isinstance(document, dict):
                metadata = document.get("metadata")
                revision = metadata.get("application_revision") if isinstance(metadata, dict) else None
            result = service.canonical_from_period_profit_export(
                document,
                scope_alias=args.scope_alias,
                revision=revision or "UNAVAILABLE",
            )
        written = _write_json(result, args.output)
        return 1 if result.get("error") or not written else 0

    try:
        reference_document = _read_json(args.reference)
        snapshot_document = _read_json(args.snapshot)
    except (OSError, ValueError, UnicodeError):
        result = service._error("DAILY_USE_NORMALIZED_INPUT_READ_FAILED")
    else:
        reference = reference_document.get("reference") if isinstance(reference_document, dict) else None
        snapshot = snapshot_document.get("snapshot") if isinstance(snapshot_document, dict) else None
        result = service.compare(reference, snapshot)
    written = _write_json(result, args.output)
    return 1 if result.get("error") or not written else 0


if __name__ == "__main__":
    sys.exit(main())
